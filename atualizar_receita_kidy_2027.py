#!/usr/bin/env python3
"""Orquestrador local, conservador, do cruzamento Receita x auditoria Kidy.
Nao atualiza a auditoria nem publica resultados no portal.
"""
import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import time
import shutil
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
AUDIT = ROOT / 'diagnosticos' / 'AUDITORIA_CNPJ_CARTEIRA_2027.xlsx'
OUT = ROOT / 'diagnosticos' / 'VALIDACAO_CNPJ_RECEITA_2027.xlsx'
STATE = ROOT / 'diagnosticos' / 'estado_atualizacao_receita_2027.json'
PROCESSOR = ROOT / 'cruzar_receita_cnpj_2027.py'


def file_sha256(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def inspect_zip(path):
    import zipfile
    with zipfile.ZipFile(path) as z:
        entries = [i for i in z.infolist() if not i.is_dir()]
        if not entries:
            raise ValueError(f'ZIP vazio: {path}')
        if any(i.flag_bits & 1 for i in entries):
            raise ValueError(f'ZIP criptografado: {path}')
        return {'name': path.name, 'size': path.stat().st_size,
                'mtime_ns': path.stat().st_mtime_ns,
                'members': [(i.filename, i.file_size, i.CRC) for i in entries]}


def valid_report(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        required = {'Resumo', 'Situacao por CNPJ', 'Vinculos Kidy', 'Pendencias CNPJ'}
        if not required.issubset(wb.sheetnames):
            raise ValueError('Relatorio incompleto: abas ausentes')
        rows = wb['Situacao por CNPJ'].iter_rows(values_only=True)
        header = next(rows, ())
        if not {'CNPJ', 'Situacao Receita', 'Resultado Cruzamento'}.issubset(set(header)):
            raise ValueError('Relatorio sem colunas esperadas')
        count = sum(1 for r in rows if r and r[0] is not None)
        if count == 0:
            raise ValueError('Relatorio sem CNPJs')
        return count
    finally:
        wb.close()


def retry_windows(operation, description, attempts=8, delay=2):
    for attempt in range(1, attempts + 1):
        try:
            return operation()
        except PermissionError as exc:
            if attempt == attempts:
                raise PermissionError(
                    f'{description}: arquivo bloqueado apos {attempts} tentativas. '
                    'Feche Excel, Explorador com visualizacao previa e sincronizadores. '
                    f'Detalhe: {exc}'
                ) from exc
            print(f'ARQUIVO OCUPADO | {description} | tentativa {attempt}/{attempts}; aguardando...', flush=True)
            time.sleep(delay)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--receita-raiz', default=r'D:\base receita federal')
    ap.add_argument('--forcar', action='store_true')
    args = ap.parse_args()
    raiz = Path(args.receita_raiz)
    if not raiz.is_dir():
        raise ValueError(f'Pasta da Receita inexistente: {raiz}')
    for p in (AUDIT, PROCESSOR):
        if not p.is_file():
            raise ValueError(f'Arquivo necessario nao encontrado: {p}')
    folders = sorted((p for p in raiz.iterdir() if p.is_dir() and re.fullmatch(r'20\d\d-(0[1-9]|1[0-2])', p.name)), reverse=True)
    if not folders:
        raise ValueError('Nenhuma pasta AAAA-MM encontrada')
    selected = None
    infos = None
    for folder in folders:
        paths = [folder / f'Estabelecimentos{i}.zip' for i in range(10)]
        if all(p.is_file() and p.stat().st_size > 0 for p in paths):
            selected, infos = folder, [inspect_zip(p) for p in paths]
            break
        print(f'IGNORANDO competencia incompleta: {folder.name}', flush=True)
    if selected is None:
        raise ValueError('Nenhuma competencia com os 10 ZIPs disponiveis')
    empresas = [selected / f'Empresas{i}.zip' for i in range(10)]
    presentes = [p.is_file() for p in empresas]
    if any(presentes) and not all(presentes):
        raise ValueError('Empresas parcialmente baixadas: conclua os 10 ZIPs ou remova os parciais')
    if all(presentes):
        infos += [inspect_zip(p) for p in empresas]
    signature = hashlib.sha256(json.dumps({
        'competencia': selected.name, 'zips': infos,
        'auditoria_sha256': file_sha256(AUDIT),
        'processador_sha256': file_sha256(PROCESSOR),
    }, sort_keys=True).encode()).hexdigest()
    prior = json.loads(STATE.read_text(encoding='utf-8')) if STATE.exists() else {}
    if not args.forcar and prior.get('signature') == signature and OUT.exists():
        try:
            n = valid_report(OUT)
            print(f'SEM ALTERACOES | competencia={selected.name} | CNPJs={n:,} | nenhum reprocessamento')
            return
        except Exception:
            print('Relatorio atual inconsistente; reprocessando', flush=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    fd, tmpname = tempfile.mkstemp(prefix='receita_2027_', suffix='.xlsx', dir=OUT.parent)
    os.close(fd)
    tmp = Path(tmpname)
    try:
        command = [sys.executable, str(PROCESSOR), '--auditoria', str(AUDIT),
                   '--receita', str(selected), '--competencia', selected.name,
                   '--saida', str(tmp)]
        print(f'PROCESSANDO | competencia={selected.name}', flush=True)
        subprocess.run(command, check=True, cwd=ROOT)
        n = valid_report(tmp)
        # Preserve a versao anterior, inclusive se a troca falhar.
        backup = OUT.with_name('VALIDACAO_CNPJ_RECEITA_2027_ANTERIOR.xlsx')
        if OUT.exists():
            retry_windows(lambda: shutil.copy2(OUT, backup), 'Backup do relatorio anterior')
        retry_windows(lambda: os.replace(tmp, OUT), 'Publicacao local do relatorio validado')
        state = {'signature': signature, 'competencia': selected.name,
                 'concluido_em': datetime.now().astimezone().isoformat(),
                 'cnpjs': n, 'auditoria_sha256': file_sha256(AUDIT)}
        temp_state = STATE.with_suffix('.json.tmp')
        temp_state.write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding='utf-8')
        retry_windows(lambda: os.replace(temp_state, STATE), 'Atualizacao do estado')
        print(f'SUCESSO | competencia={selected.name} | CNPJs={n:,} | arquivo={OUT}', flush=True)
        print('NAO PUBLICADO: Drive, Supabase e Vercel permanecem inalterados', flush=True)
    finally:
        if tmp.exists():
            try:
                retry_windows(lambda: tmp.unlink(), 'Limpeza do temporario', attempts=3)
            except PermissionError as exc:
                print(f'AVISO: temporario mantido para limpeza posterior: {tmp} ({exc})', flush=True)

if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        print(f'ERRO: {exc}', file=sys.stderr, flush=True)
        sys.exit(1)
