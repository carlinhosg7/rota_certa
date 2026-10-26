#!/usr/bin/env python3
"""Cruza auditoria Kidy com ZIPs oficiais de Dados Abertos CNPJ da RFB.
Execução local; não altera fonte, Drive ou portal.
"""
import argparse
import csv
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

SITUACOES = {'01':'NULA','02':'ATIVA','03':'SUSPENSA','04':'INAPTA','08':'BAIXADA'}

def so_digitos(v):
    return re.sub(r'\D', '', str(v or '').strip())

def check_cnpj(s):
    if len(s) != 14 or len(set(s)) == 1 or not s.isdigit(): return False
    for n, pesos in [(12,[5,4,3,2,9,8,7,6,5,4,3,2]),(13,[6,5,4,3,2,9,8,7,6,5,4,3,2])]:
        resto = sum(int(a)*b for a,b in zip(s[:n], pesos))%11
        if int(s[n]) != (0 if resto < 2 else 11-resto): return False
    return True

def linhas_zip(caminho):
    with zipfile.ZipFile(caminho) as z:
        for nome in z.namelist():
            if nome.endswith('/') or nome.lower().endswith(('.pdf','.txt')): continue
            with z.open(nome) as arq:
                import io
                stream = io.TextIOWrapper(arq, encoding='latin-1', errors='replace', newline='')
                yield from csv.reader(stream, delimiter=';', quotechar='"')

def escrever(ws, valores):
    ws.append([str(v) if v is not None else '' for v in valores])

def main():
    p = argparse.ArgumentParser(description='Cruzamento local de CNPJ com dados abertos da Receita Federal')
    p.add_argument('--auditoria', required=True, type=Path)
    p.add_argument('--receita', required=True, type=Path, help='Pasta contendo Estabelecimentos0.zip a Estabelecimentos9.zip e Empresas0.zip a Empresas9.zip')
    p.add_argument('--competencia', required=True, help='AAAA-MM da base oficial baixada')
    p.add_argument('--saida', required=True, type=Path)
    args = p.parse_args()
    if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', args.competencia): p.error('Competência deve ser AAAA-MM')
    if not args.auditoria.is_file(): p.error('Arquivo de auditoria não encontrado')
    estabs = sorted(args.receita.glob('Estabelecimentos*.zip'))
    empresas = sorted(args.receita.glob('Empresas*.zip'))
    def partes_esperadas(arquivos, prefixo):
        achados = {x.name.lower() for x in arquivos}
        faltando = [f'{prefixo}{i}.zip' for i in range(10) if f'{prefixo}{i}.zip'.lower() not in achados]
        if faltando: p.error('Arquivos faltando: ' + ', '.join(faltando))
    partes_esperadas(estabs,'Estabelecimentos')
    if empresas:
        partes_esperadas(empresas,'Empresas')
    else:
        print('AVISO: arquivos Empresas ausentes; razao social oficial ficara vazia.', flush=True)

    wb = load_workbook(args.auditoria, read_only=True, data_only=True)
    candidatos = defaultdict(list)
    pendencias = []
    total_linhas = 0
    for aba in ('Validos','Recuperaveis','Invalidos','Sem CNPJ'):
        if aba not in wb.sheetnames: p.error(f'Aba ausente na auditoria: {aba}')
        rows = wb[aba].iter_rows(values_only=True)
        cab = next(rows, None)
        if not cab: continue
        col = {str(v).strip(): i for i,v in enumerate(cab) if v is not None}
        for requerido in ('CNPJ Original','CNPJ Normalizado','Categoria','Codigo Cliente'):
            if requerido not in col: p.error(f'Coluna ausente em {aba}: {requerido}')
        for row in rows:
            total_linhas += 1
            get = lambda k: str(row[col[k]] or '').strip() if k in col and col[k] < len(row) else ''
            dado = {'codigo':get('Codigo Cliente'),'razao_kidy':get('Razao Social'),'original':get('CNPJ Original'),
                    'cnpj':so_digitos(get('CNPJ Normalizado')),'grupo':get('Grupo Cliente'),
                    'representante':get('Representante'),'supervisor':get('Supervisor'),
                    'categoria':get('Categoria'),'motivo':get('Motivo')}
            if dado['categoria'] in ('VALIDO','RECUPERAVEL') and check_cnpj(dado['cnpj']):
                candidatos[dado['cnpj']].append(dado)
            else: pendencias.append(dado)
    wb.close()
    print(f'CNPJs candidatos distintos após normalização: {len(candidatos):,}', flush=True)
    print(f'Linhas na auditoria: {total_linhas:,} (a auditoria anterior mantém somente exemplos por CNPJ original)',flush=True)
    print(f'Pendências de formato: {len(pendencias):,}',flush=True)

    encontrados = {}
    raizes = set()
    for i, arquivo in enumerate(estabs, 1):
        print(f'[{i}/10] Estabelecimentos: {arquivo.name}', flush=True)
        for r in linhas_zip(arquivo):
            if len(r) < 20: continue
            cnpj = ''.join(r[:3]).strip()
            if cnpj not in candidatos: continue
            if cnpj in encontrados: continue
            raizes.add(cnpj[:8])
            encontrados[cnpj] = {
                'situacao_codigo':r[5].strip(), 'situacao_data':r[6].strip(), 'motivo_codigo':r[7].strip(),
                'abertura':r[10].strip(), 'cnae':r[11].strip(), 'uf':r[19].strip(),
                'matriz_filial':r[3].strip(), 'fantasia':r[4].strip(), 'razao_receita':''
            }
    print(f'CNPJs encontrados nos estabelecimentos: {len(encontrados):,}',flush=True)
    empresas_por_raiz = {}
    for i, arquivo in enumerate(empresas, 1):
        print(f'[{i}/10] Empresas: {arquivo.name}', flush=True)
        for r in linhas_zip(arquivo):
            if len(r) >= 2 and r[0].strip() in raizes:
                empresas_por_raiz[r[0].strip()] = r[1].strip()
    for cnpj, d in encontrados.items(): d['razao_receita'] = empresas_por_raiz.get(cnpj[:8], '')

    args.saida.parent.mkdir(parents=True, exist_ok=True)
    out = Workbook(write_only=True)
    resumo = out.create_sheet('Resumo')
    escrever(resumo, ['Indicador','Quantidade'])
    escrever(resumo, ['Competencia Receita',args.competencia])
    escrever(resumo, ['CNPJs candidatos distintos',len(candidatos)])
    escrever(resumo, ['CNPJs encontrados',len(encontrados)])
    escrever(resumo, ['CNPJs nao localizados',len(candidatos)-len(encontrados)])
    escrever(resumo, ['Pendencias de formato (linhas)',len(pendencias)])
    cont = Counter(SITUACOES.get(v['situacao_codigo'],'CODIGO '+v['situacao_codigo']) for v in encontrados.values())
    for k,v in sorted(cont.items()): escrever(resumo,['Receita - '+k,v])

    cnpj_ws = out.create_sheet('Situacao por CNPJ')
    colunas = ['CNPJ','Situacao Receita','Codigo Situacao','Data Situacao','Motivo Codigo','Razao Social Receita',
               'Nome Fantasia','Data Abertura','CNAE Principal','UF','Matriz Filial','Competencia Receita',
               'Origem CNPJ','Confianca Normalizacao','Resultado Cruzamento']
    escrever(cnpj_ws,colunas)
    for cnpj in sorted(candidatos):
        d = encontrados.get(cnpj)
        origem = 'RECUPERAVEL' if any(x['categoria']=='RECUPERAVEL' for x in candidatos[cnpj]) else 'VALIDO'
        if d:
            escrever(cnpj_ws,[cnpj,SITUACOES.get(d['situacao_codigo'],'DESCONHECIDA'),d['situacao_codigo'],d['situacao_data'],
                d['motivo_codigo'],d['razao_receita'],d['fantasia'],d['abertura'],d['cnae'],d['uf'],d['matriz_filial'],
                args.competencia,origem,'CONFIRMADO NA RECEITA' if origem=='RECUPERAVEL' else 'ORIGINAL VALIDO','LOCALIZADO'])
        else:
            escrever(cnpj_ws,[cnpj,'NAO LOCALIZADO','','','','','','','','','',args.competencia,origem,
                'RECUPERACAO NAO CONFIRMADA' if origem=='RECUPERAVEL' else 'SEM CORRESPONDENCIA','NAO LOCALIZADO'])

    vinc = out.create_sheet('Vinculos Kidy')
    escrever(vinc,['Codigo Cliente','Razao Social Kidy','CNPJ Original','CNPJ Normalizado','Grupo Cliente',
                  'Representante','Supervisor','Origem','Situacao Receita','Resultado Cruzamento','Competencia Receita'])
    for cnpj, itens in candidatos.items():
        situacao = SITUACOES.get(encontrados[cnpj]['situacao_codigo'],'DESCONHECIDA') if cnpj in encontrados else 'NAO LOCALIZADO'
        for d in itens:
            escrever(vinc,[d['codigo'],d['razao_kidy'],d['original'],cnpj,d['grupo'],d['representante'],d['supervisor'],
                           d['categoria'],situacao,'LOCALIZADO' if cnpj in encontrados else 'NAO LOCALIZADO',args.competencia])
    pend = out.create_sheet('Pendencias CNPJ')
    escrever(pend,['Codigo Cliente','Razao Social Kidy','CNPJ Original','CNPJ Normalizado','Grupo Cliente',
                  'Representante','Supervisor','Categoria','Motivo'])
    for d in pendencias:
        escrever(pend,[d['codigo'],d['razao_kidy'],d['original'],d['cnpj'],d['grupo'],d['representante'],d['supervisor'],d['categoria'],d['motivo']])
    out.save(args.saida)
    print('RELATORIO GERADO:',args.saida,flush=True)
    print('Nenhuma alteracao no Google Drive, Supabase, Vercel ou codigo comercial.',flush=True)

if __name__ == '__main__':
    try: main()
    except (OSError, zipfile.BadZipFile, csv.Error) as exc:
        print('ERRO:',exc,file=sys.stderr)
        sys.exit(1)
