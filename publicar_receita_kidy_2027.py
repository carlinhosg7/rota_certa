#!/usr/bin/env python3
"""Publicacao versionada da dimensao CNPJ Kidy no Supabase. Python 3.10+, openpyxl.
Uso: python publicar_receita_kidy_2027.py --validar ; python publicar_receita_kidy_2027.py --publicar
"""
import argparse, hashlib, json, os, re, sys, time, uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
REPORT = ROOT / 'diagnosticos' / 'VALIDACAO_CNPJ_RECEITA_2027.xlsx'
FIELDS = {'CNPJ':'cnpj','Situacao Receita':'situacao_receita','Codigo Situacao':'codigo_situacao',
 'Data Situacao':'data_situacao','Motivo Codigo':'motivo_codigo','Razao Social Receita':'razao_social_receita',
 'Nome Fantasia':'nome_fantasia','Data Abertura':'data_abertura','CNAE Principal':'cnae_principal',
 'UF':'uf','Matriz Filial':'matriz_filial','Origem CNPJ':'origem_cnpj',
 'Confianca Normalizacao':'confianca_normalizacao','Resultado Cruzamento':'resultado_cruzamento'}
VALID_STATUS={'ATIVA','BAIXADA','INAPTA','SUSPENSA','NULA','NAO LOCALIZADO'}

def env_local():
    p=ROOT/'.env.local'
    if not p.exists(): return
    for line in p.read_text(encoding='utf-8-sig').splitlines():
        line=line.strip()
        if not line or line.startswith('#'): continue
        if line.startswith('export '): line=line[7:]
        if '=' not in line: continue
        k,v=line.split('=',1); k=k.strip(); v=v.strip().strip('"').strip("'")
        os.environ.setdefault(k,v)

def val(v):
    if v is None: return None
    s=str(v).strip()
    return s or None

def dte(v):
    if v is None or v=='': return None
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    s=str(v).strip()
    if not s or s in ('0','00000000'): return None
    if re.fullmatch(r'\d{8}',s):
        try: return datetime.strptime(s,'%Y%m%d').date().isoformat()
        except ValueError: return None
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}',s):
        try: return date.fromisoformat(s).isoformat()
        except ValueError: return None
    if re.fullmatch(r'\d{2}/\d{2}/\d{4}',s):
        try: return datetime.strptime(s,'%d/%m/%Y').date().isoformat()
        except ValueError: return None
    raise ValueError(f'Data nao reconhecida: {s!r}')

def ler(path):
    if not path.exists(): raise FileNotFoundError(path)
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb['Situacao por CNPJ']; it=ws.iter_rows(values_only=True)
        heads=[str(x or '').strip() for x in next(it)]
        missing=[x for x in FIELDS if x not in heads]
        if missing: raise ValueError('Colunas ausentes: '+', '.join(missing))
        idx={x:heads.index(x) for x in FIELDS}
        competencia_idx=heads.index('Competencia Receita')
        result=[]; seen=set(); comps=set(); counter=Counter(); localizados=0
        for excel_line, row in enumerate(it,2):
            if not row or not any(x is not None for x in row): continue
            r={dest:val(row[idx[src]]) if idx[src]<len(row) else None for src,dest in FIELDS.items()}
            cnpj=re.sub(r'\D','',r['cnpj'] or '')
            if len(cnpj)!=14: raise ValueError(f'CNPJ invalido na linha {excel_line}: {r["cnpj"]!r}')
            if cnpj in seen: raise ValueError(f'CNPJ duplicado: {cnpj} na linha {excel_line}')
            seen.add(cnpj); r['cnpj']=cnpj
            for k in ('data_situacao','data_abertura'): r[k]=dte(r[k])
            situacao=(r['situacao_receita'] or '').upper().strip()
            if situacao not in VALID_STATUS: raise ValueError(f'Situacao inesperada {situacao!r}, linha {excel_line}')
            r['situacao_receita']=situacao
            r['codigo_situacao']=str(r['codigo_situacao']).zfill(2) if r['codigo_situacao'] else None
            r['uf']=(r['uf'] or '').upper() or None
            if r['uf'] and len(r['uf'])!=2: raise ValueError(f'UF invalida linha {excel_line}')
            r['resultado_cruzamento']=(r['resultado_cruzamento'] or '').upper()
            if r['resultado_cruzamento']=='LOCALIZADO': localizados+=1
            comp=val(row[competencia_idx]) if competencia_idx<len(row) else None
            if comp: comps.add(comp)
            counter[situacao]+=1; result.append(r)
        if len(comps)!=1: raise ValueError(f'Competencias inconsistentes: {sorted(comps)}')
        competencia=comps.pop()
        if not re.fullmatch(r'20\d\d-(0[1-9]|1[0-2])',competencia): raise ValueError('Competencia invalida')
        if len(result)<1: raise ValueError('Relatorio vazio')
        if localizados!=sum(n for s,n in counter.items() if s!='NAO LOCALIZADO'):
            raise ValueError('Divergencia entre situacao e resultado de cruzamento')
        return result,competencia,localizados,counter
    finally: wb.close()

def req(method, url, key, path, payload=None, extra=None, attempts=4):
    endpoint=url.rstrip('/')+'/rest/v1/'+path
    body=json.dumps(payload,ensure_ascii=False).encode('utf-8') if payload is not None else None
    headers={'apikey':key,'Authorization':'Bearer '+key,'Content-Type':'application/json','Accept':'application/json'}
    if extra: headers.update(extra)
    for attempt in range(attempts):
        try:
            with urlopen(Request(endpoint,data=body,headers=headers,method=method),timeout=120) as response:
                content=response.read().decode('utf-8')
                return json.loads(content) if content else None,dict(response.headers)
        except HTTPError as e:
            err=e.read().decode('utf-8','replace')[:1200]
            if e.code in (429,500,502,503,504) and attempt+1<attempts:
                time.sleep(2**attempt);continue
            raise RuntimeError(f'Supabase HTTP {e.code} em {path}: {err}') from e
        except (URLError,TimeoutError) as e:
            if attempt+1<attempts: time.sleep(2**attempt);continue
            raise RuntimeError(f'Falha de rede em {path}: {e}') from e

def main():
    ap=argparse.ArgumentParser()
    group=ap.add_mutually_exclusive_group(required=True)
    group.add_argument('--validar',action='store_true',help='Somente verifica o Excel local')
    group.add_argument('--publicar',action='store_true',help='Envia, valida e ativa a versao no Supabase')
    ap.add_argument('--arquivo',type=Path,default=REPORT)
    args=ap.parse_args()
    registros,comp,localizados,contagem=ler(args.arquivo)
    digest=hashlib.sha256(args.arquivo.read_bytes()).hexdigest()
    print(f'VALIDADO | competencia={comp} | CNPJs={len(registros):,} | localizados={localizados:,}',flush=True)
    print('SITUACOES | '+json.dumps(contagem,ensure_ascii=False),flush=True)
    if args.validar: return
    env_local()
    url=os.environ.get('NEXT_PUBLIC_SUPABASE_URL') or os.environ.get('SUPABASE_URL')
    key=os.environ.get('SUPABASE_SECRET_KEY') or os.environ.get('SUPABASE_SERVICE_ROLE_KEY')
    if not url or not key: raise RuntimeError('URL ou chave secreta do Supabase ausente no ambiente')
    # Não imprimir chave ou URL.
    existing,_=req('GET',url,key,'kidy_receita_versoes?select=id,situacao&origem_sha256=eq.'+digest+'&limit=1')
    if existing:
        status=existing[0]['situacao']
        if status=='ATIVA':
            print('SEM ALTERACOES | mesma origem ja ATIVA; nenhuma publicacao');return
        if status!='CARREGANDO':
            raise RuntimeError(f'Hash ja existe em estado {status}; verificar antes de republicar')
        versao=existing[0]['id'];print('RETOMANDO | versao ja criada',flush=True)
    else:
        versao=str(uuid.uuid4())
        req('POST',url,key,'kidy_receita_versoes', [{
            'id':versao,'competencia':comp,'origem_sha256':digest,'situacao':'CARREGANDO',
            'total_cnpjs':len(registros),'total_localizados':localizados,
            'observacao':'Carga automatizada do relatorio local validado'
        }],{'Prefer':'return=minimal'})
        print('VERSAO CRIADA | estado=CARREGANDO',flush=True)
    lote=300
    for start in range(0,len(registros),lote):
        chunk=[{'versao_id':versao,**r} for r in registros[start:start+lote]]
        # UPSERT idempotente exige PK (versao_id, cnpj) conforme schema.
        req('POST',url,key,'kidy_receita_cnpj?on_conflict=versao_id,cnpj',chunk,
            {'Prefer':'resolution=merge-duplicates,return=minimal'})
        if (start//lote+1)%10==0 or start+lote>=len(registros):
            print(f'ENVIADOS | {min(start+lote,len(registros)):,}/{len(registros):,}',flush=True)
    # Funcao SQL valida contagem, distribuicao e troca ativa atomicamente.
    result,_=req('POST',url,key,'rpc/kidy_receita_ativar_versao',{
        'p_versao_id':versao,'p_total':len(registros),'p_localizados':localizados,
        'p_distribuicao':dict(contagem)
    })
    print('PUBLICACAO CONCLUIDA | versao ativa | '+json.dumps(result,ensure_ascii=False),flush=True)

if __name__=='__main__':
    try: main()
    except Exception as e:
        print('ERRO | '+str(e),file=sys.stderr);sys.exit(1)
