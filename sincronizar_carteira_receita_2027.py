#!/usr/bin/env python3
"""Sincroniza fonte Drive, auditoria, Receita e vinculos completos, apenas localmente."""
import hashlib, json, os, subprocess, sys, tempfile, shutil, re, time
from pathlib import Path
from collections import Counter
from openpyxl import Workbook, load_workbook

ROOT=Path(__file__).resolve().parent
DIAG=ROOT/'diagnosticos'
SOURCE=DIAG/'FONTE_CARTEIRA_2027.xlsx'
AUDIT=DIAG/'AUDITORIA_CNPJ_CARTEIRA_2027.xlsx'
REPORT=DIAG/'VALIDACAO_CNPJ_RECEITA_2027.xlsx'
OUTPUT=DIAG/'CARTEIRA_2027_RECEITA_CONSOLIDADA.xlsx'
STATE=DIAG/'estado_sincronizacao_carteira_2027.json'
NAME='DADOS PRETIVIA CLIENTES.xlsx'
FIELDS=['Codigo Cliente','Razao Social Empresa','Cnpj','Codigo Grupo Cliente','Grupo Cliente','Codigo Representante','Representante','Codigo Supervisor','Supervisor','Bloqueado Financeiro','Cliente Ativo','Limite Aprovado','Limite Disponivel','Data Ultima Compra','Qtd Venda','Vlr Venda']
AUD_FIELDS=['Codigo Cliente','Razao Social','CNPJ Original','CNPJ Normalizado','Categoria','Grupo Cliente','Representante','Supervisor','Motivo']

def retry_windows(action, description, attempts=8, delay=2):
 for attempt in range(1, attempts+1):
  try:return action()
  except PermissionError:
   if attempt == attempts:raise
   print(f'ARQUIVO OCUPADO | {description} | tentativa {attempt}/{attempts}',flush=True)
   time.sleep(delay)

def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
 return h.hexdigest()

def digits(v):
 if v is None:return ''
 if isinstance(v,int):return str(v)
 if isinstance(v,float):return str(int(v)) if v.is_integer() else re.sub(r'\D','',str(v))
 return re.sub(r'\D','',str(v))

def valid(s):
 if len(s)!=14 or len(set(s))==1 or not s.isdigit():return False
 for n,w in ((12,[5,4,3,2,9,8,7,6,5,4,3,2]),(13,[6,5,4,3,2,9,8,7,6,5,4,3,2])):
  rem=sum(int(a)*b for a,b in zip(s[:n],w))%11
  if int(s[n])!=(0 if rem<2 else 11-rem):return False
 return True

def classify(value):
 s=digits(value)
 if not s:return ('SEM CNPJ','','CNPJ ausente')
 if valid(s):return ('VALIDO',s,'CNPJ valido')
 if len(s)<14 and valid(s.zfill(14)):return ('RECUPERAVEL',s.zfill(14),'Zeros a esquerda; validar identidade')
 return ('INVALIDO','','Formato ou digitos verificadores invalidos')

def rows_source(p):
 wb=load_workbook(p,read_only=True,data_only=True)
 try:
  ws=wb[NAME.replace('.xlsx','')] if NAME.replace('.xlsx','') in wb.sheetnames else wb.worksheets[0]
  rows=ws.iter_rows(values_only=True)
  for i,row in enumerate(rows,1):
   labels=[str(x or '').strip().lower() for x in row]
   if 'codigo cliente' in labels and 'cnpj' in labels:
    idx={str(v or '').strip().lower():j for j,v in enumerate(row)}
    missing=[f for f in FIELDS if f.lower() not in idx]
    if missing:raise RuntimeError('Colunas ausentes: '+', '.join(missing))
    break
  else:raise RuntimeError('Cabecalho da base nao localizado')
  for row in rows:
   if not any(v is not None for v in row):continue
   yield {f:(row[idx[f.lower()]] if idx[f.lower()]<len(row) else None) for f in FIELDS}
 finally:wb.close()

def download():
 js=r'''const fs=require('fs');const path=require('path');
const {loadEnvConfig}=require('@next/env');loadEnvConfig(process.cwd());
const {listarArquivos,baixarArquivo}=require('./src/lib/google-drive-data.cjs');
(async()=>{const f=(await listarArquivos()).find(x=>x.nome==='DADOS PRETIVIA CLIENTES.xlsx');if(!f)throw Error('Fonte nao encontrada no Drive');
const b=await baixarArquivo(f);fs.writeFileSync(process.argv[1],b);console.log('DRIVE_OK | bytes='+b.length)})().catch(e=>{console.error('DRIVE_ERRO: '+e.message);process.exitCode=1});'''
 fd,tmp=tempfile.mkstemp(prefix='drive_kidy_',suffix='.xlsx',dir=DIAG);os.close(fd)
 try:
  subprocess.run(['node','-e',js,tmp],cwd=ROOT,check=True)
  if not zip_check(Path(tmp)):raise RuntimeError('Arquivo baixado nao e um XLSX valido')
  return Path(tmp)
 except:
  Path(tmp).unlink(missing_ok=True);raise

def zip_check(p):
 import zipfile
 return zipfile.is_zipfile(p)

def build_audit(src, dest):
 w=Workbook(write_only=True)
 sheets={k:w.create_sheet(k) for k in ('Validos','Recuperaveis','Invalidos','Sem CNPJ')}
 for ws in sheets.values():ws.append(AUD_FIELDS)
 counts=Counter()
 for r in rows_source(src):
  cat,cnpj,reason=classify(r['Cnpj']);counts[cat]+=1
  sheet={'VALIDO':'Validos','RECUPERAVEL':'Recuperaveis','INVALIDO':'Invalidos','SEM CNPJ':'Sem CNPJ'}[cat]
  sheets[sheet].append([r['Codigo Cliente'],r['Razao Social Empresa'],str(r['Cnpj'] or ''),cnpj,cat,r['Grupo Cliente'],r['Representante'],r['Supervisor'],reason])
 w.save(dest)
 print('AUDITORIA | '+str(dict(counts))+' | registros='+str(sum(counts.values())),flush=True)
 return sum(counts.values())

def receipt_lookup():
 w=load_workbook(REPORT,read_only=True,data_only=True)
 try:
  ws=w['Situacao por CNPJ'];rows=ws.iter_rows(values_only=True);head=next(rows)
  idx={str(v).strip():i for i,v in enumerate(head)}
  keys=['CNPJ','Situacao Receita','Codigo Situacao','Data Situacao','Razao Social Receita','Competencia Receita','Resultado Cruzamento']
  for k in keys:
   if k not in idx:raise RuntimeError('Coluna ausente no relatorio Receita: '+k)
  result={}
  for row in rows:
   c=digits(row[idx['CNPJ']]);
   if len(c)==14:result[c]={k:row[idx[k]] for k in keys[1:]}
  return result
 finally:w.close()

def build_join(src,dest):
 lookup=receipt_lookup()
 w=Workbook(write_only=True);ws=w.create_sheet('Carteira Completa')
 extra=['CNPJ Original','CNPJ Normalizado','Qualidade CNPJ','Situacao Receita','Codigo Situacao','Data Situacao','Razao Social Receita','Competencia Receita','Resultado Cruzamento']
 ws.append(FIELDS+extra)
 counts=Counter();total=0
 for r in rows_source(src):
  cat,c,reason=classify(r['Cnpj']);rec=lookup.get(c,{})
  status=rec.get('Situacao Receita','PENDENTE FORMATO' if not c else 'NAO LOCALIZADO')
  ws.append([r[f] for f in FIELDS]+[str(r['Cnpj'] or ''),c,cat,status,rec.get('Codigo Situacao',''),rec.get('Data Situacao',''),rec.get('Razao Social Receita',''),rec.get('Competencia Receita',''),rec.get('Resultado Cruzamento','NAO APLICAVEL' if not c else 'NAO LOCALIZADO')])
  counts[str(status)]+=1;total+=1
 sumws=w.create_sheet('Resumo');sumws.append(['Indicador','Quantidade']);sumws.append(['Registros comerciais preservados',total]);sumws.append(['CNPJs oficiais distintos no relatorio',len(lookup)])
 for k,v in sorted(counts.items()):sumws.append([k,v])
 w.save(dest)
 print('CONSOLIDADO | registros='+str(total)+' | '+str(dict(counts)),flush=True)
 return total

def main():
 DIAG.mkdir(parents=True,exist_ok=True)
 if not (ROOT/'atualizar_receita_kidy_2027.py').exists():raise RuntimeError('Atualizador Receita nao encontrado')
 downloaded=download()
 try:
  current=sha(downloaded)
  prior=json.loads(STATE.read_text(encoding='utf-8')) if STATE.exists() else {}
  changed=not SOURCE.exists() or sha(SOURCE)!=current
  if changed:
   # Validar em OUTRO processo: no Windows, openpyxl pode manter
   # um ZipExtFile aberto mesmo depois de wb.close(). Ao terminar o
   # subprocesso, o Windows libera todos os handles do XLSX.
   validator = """import sys
from openpyxl import load_workbook
p=sys.argv[1]
w=load_workbook(p,read_only=True,data_only=True)
try:
 ok=False
 for ws in w.worksheets:
  for i,row in enumerate(ws.iter_rows(values_only=True),1):
   names={str(v or '').strip().lower() for v in row}
   if 'codigo cliente' in names and 'cnpj' in names:
    ok=any(any(v is not None for v in r) for r in ws.iter_rows(min_row=i+1,values_only=True))
    if ok:break
  if ok:break
 if not ok:raise RuntimeError('Base Drive sem cabecalho ou clientes')
finally:w.close()
"""
   subprocess.run([sys.executable,'-c',validator,str(downloaded)],check=True,cwd=ROOT)
   backup=SOURCE.with_name('FONTE_CARTEIRA_2027_ANTERIOR.xlsx')
   if SOURCE.exists():shutil.copy2(SOURCE,backup)
   retry_windows(lambda:os.replace(downloaded,SOURCE),'Substituir fonte Drive')
   print('FONTE | atualizada a partir do Drive',flush=True)
  else:print('FONTE | sem alteracoes',flush=True)
 finally:downloaded.unlink(missing_ok=True)
 if changed or not AUDIT.exists():
  fd,tmp=tempfile.mkstemp(prefix='audit_',suffix='.xlsx',dir=DIAG);os.close(fd)
  try:
   build_audit(SOURCE,Path(tmp))
   retry_windows(lambda:os.replace(tmp,AUDIT),'Atualizar auditoria')
  finally:Path(tmp).unlink(missing_ok=True)
 else:print('AUDITORIA | sem alteracoes',flush=True)
 subprocess.run([sys.executable,str(ROOT/'atualizar_receita_kidy_2027.py')],cwd=ROOT,check=True)
 sig={'source':sha(SOURCE),'receipt':sha(REPORT),'script':sha(Path(__file__))}
 if prior.get('signature')==sig and OUTPUT.exists():
  print('SEM ALTERACOES | carteira consolidada ja atualizada');return
 fd,tmp=tempfile.mkstemp(prefix='carteira_',suffix='.xlsx',dir=DIAG);os.close(fd)
 try:
  n=build_join(SOURCE,Path(tmp))
  if OUTPUT.exists():shutil.copy2(OUTPUT,OUTPUT.with_name('CARTEIRA_2027_RECEITA_CONSOLIDADA_ANTERIOR.xlsx'))
  retry_windows(lambda:os.replace(tmp,OUTPUT),'Atualizar carteira consolidada')
 finally:Path(tmp).unlink(missing_ok=True)
 state_tmp=STATE.with_suffix('.json.tmp');state_tmp.write_text(json.dumps({'signature':sig,'registros':n},indent=2),encoding='utf-8');retry_windows(lambda:os.replace(state_tmp,STATE),'Atualizar estado')
 print('SUCESSO | '+str(OUTPUT));print('LOCAL SOMENTE | nada publicado no Drive, Supabase ou Vercel')

if __name__=='__main__':
 try:main()
 except Exception as e:print('ERRO: '+str(e),file=sys.stderr);sys.exit(1)
